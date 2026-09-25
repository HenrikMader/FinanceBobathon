"""
Helper script — run once to (re-)generate all lab data artefacts.
Not part of the lab itself; participants never see or run this.
"""
import csv
import random
from datetime import date, timedelta
from pathlib import Path

# ── optional: openpyxl for Excel ────────────────────────────────────────────
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("openpyxl not installed — skipping Excel file creation.")

DATA_DIR = Path(__file__).parent.parent / "data"
DATA_DIR.mkdir(exist_ok=True)

random.seed(42)

# ── 1. budget_tracker.xlsx ───────────────────────────────────────────────────

COST_CENTERS = [
    "CC100 — Sales",
    "CC200 — Marketing",
    "CC300 — Operations",
    "CC400 — IT",
    "CC500 — Finance",
]
QUARTERS = ["Q1", "Q2", "Q3", "Q4"]

# Plan budgets (annual, then split per quarter slightly unevenly)
ANNUAL_PLAN = {
    "CC100 — Sales":      1_200_000,
    "CC200 — Marketing":    800_000,
    "CC300 — Operations": 2_400_000,
    "CC400 — IT":           600_000,
    "CC500 — Finance":      400_000,
}

Q_SPLIT = {  # fractions per quarter (must sum to 1)
    "Q1": 0.20,
    "Q2": 0.25,
    "Q3": 0.30,
    "Q4": 0.25,
}

# Actual spend — inject a few interesting variances
VARIANCE_OVERRIDES = {
    # (cc, quarter): multiplier on plan
    ("CC200 — Marketing", "Q1"): 1.35,   # over budget — campaign spike
    ("CC300 — Operations", "Q3"): 0.78,  # under budget — vacancy savings
    ("CC400 — IT", "Q2"): 1.18,          # over budget — unplanned licence renewal
    ("CC100 — Sales", "Q4"): 0.92,       # slightly under
}


def make_budget_xlsx():
    if not HAS_OPENPYXL:
        return

    wb = openpyxl.Workbook()

    # ── Sheet 1: Budget Overview ─────────────────────────────────────────────
    ws = wb.active
    ws.title = "Budget Overview"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    sub_fill    = PatternFill("solid", fgColor="D6E4F0")
    total_fill  = PatternFill("solid", fgColor="BDD7EE")
    over_fill   = PatternFill("solid", fgColor="FFCCCC")
    thin        = Side(style="thin", color="AAAAAA")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title row
    ws.merge_cells("A1:L1")
    ws["A1"] = "Acme GmbH — Budget Controlling 2024 (Plan vs. Ist)"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")

    # Column headers — row 3
    headers = ["Kostenstelle"]
    for q in QUARTERS:
        headers += [f"{q} Plan (€)", f"{q} Ist (€)", f"{q} Abw. (€)"]
    headers += ["Jahres-Plan (€)", "Jahres-Ist (€)", "Jahres-Abw. (€)", "Abw. %"]

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = border

    ws.row_dimensions[3].height = 30
    ws.column_dimensions["A"].width = 24
    for i in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14

    # Data rows
    row = 4
    for cc in COST_CENTERS:
        plan_annual = ANNUAL_PLAN[cc]
        ws.cell(row=row, column=1, value=cc).font = Font(bold=True)

        col = 2
        plan_cells = []
        ist_cells  = []

        for q in QUARTERS:
            plan_val = round(plan_annual * Q_SPLIT[q])
            mult = VARIANCE_OVERRIDES.get((cc, q), random.uniform(0.93, 1.07))
            ist_val  = round(plan_val * mult)

            plan_col_letter = get_column_letter(col)
            ist_col_letter  = get_column_letter(col + 1)
            abw_col_letter  = get_column_letter(col + 2)

            p_cell = ws.cell(row=row, column=col,     value=plan_val)
            i_cell = ws.cell(row=row, column=col + 1, value=ist_val)
            a_cell = ws.cell(row=row, column=col + 2,
                             value=f"={ist_col_letter}{row}-{plan_col_letter}{row}")

            for c in [p_cell, i_cell, a_cell]:
                c.number_format = '#,##0'
                c.border = border
                c.alignment = Alignment(horizontal="right")

            # Highlight over-budget deviation
            if ist_val > plan_val * 1.10:
                a_cell.fill = over_fill

            plan_cells.append(plan_col_letter)
            ist_cells.append(ist_col_letter)
            col += 3

        # Annual totals
        plan_sum_formula = "=" + "+".join(f"{l}{row}" for l in plan_cells)
        ist_sum_formula  = "=" + "+".join(f"{l}{row}" for l in ist_cells)

        plan_total_col = col
        ist_total_col  = col + 1
        abw_total_col  = col + 2
        pct_col        = col + 3

        pt = ws.cell(row=row, column=plan_total_col, value=plan_sum_formula)
        it = ws.cell(row=row, column=ist_total_col,  value=ist_sum_formula)
        at = ws.cell(row=row, column=abw_total_col,
                     value=f"={get_column_letter(ist_total_col)}{row}"
                           f"-{get_column_letter(plan_total_col)}{row}")

        # BUG deliberately introduced: wrong formula — divides by Ist instead of Plan
        # Participants will discover and fix this with Bob's help
        pct = ws.cell(row=row, column=pct_col,
                      value=f"={get_column_letter(abw_total_col)}{row}"
                            f"/{get_column_letter(ist_total_col)}{row}")  # <-- BUG: should be plan_total_col

        for c in [pt, it, at]:
            c.number_format = '#,##0'
            c.fill = total_fill
            c.border = border
            c.alignment = Alignment(horizontal="right")
            c.font = Font(bold=True)

        pct.number_format = '0.0%'
        pct.fill = total_fill
        pct.border = border
        pct.alignment = Alignment(horizontal="right")
        pct.font = Font(bold=True)

        row += 1

    # Grand total row
    ws.cell(row=row, column=1, value="GESAMT").font = Font(bold=True, size=11)
    col = 2
    for _ in QUARTERS:
        for offset in range(3):
            c = ws.cell(row=row, column=col + offset,
                        value=f"=SUM({get_column_letter(col + offset)}4"
                              f":{get_column_letter(col + offset)}{row - 1})")
            c.number_format = '#,##0'
            c.fill = PatternFill("solid", fgColor="1F4E79")
            c.font = Font(bold=True, color="FFFFFF")
            c.border = border
            c.alignment = Alignment(horizontal="right")
        col += 3

    for offset in range(4):
        c = ws.cell(row=row, column=col + offset,
                    value=f"=SUM({get_column_letter(col + offset)}4"
                          f":{get_column_letter(col + offset)}{row - 1})")
        c.number_format = '#,##0' if offset < 3 else '0.0%'
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.font = Font(bold=True, color="FFFFFF")
        c.border = border
        c.alignment = Alignment(horizontal="right")

    ws.freeze_panes = "B4"

    # ── Sheet 2: Notes (deliberately sparse — Bob will explain) ─────────────
    ws2 = wb.create_sheet("Notes & Assumptions")
    ws2["A1"] = "Assumptions"
    ws2["A1"].font = Font(bold=True, size=12)
    notes = [
        ("Budget Year:", "2024"),
        ("Currency:", "EUR"),
        ("Plan Source:", "Annual Budget approved by CFO, Nov 2023"),
        ("Actuals Source:", "SAP FI export — last updated 2024-12-01"),
        ("Note:", "Q3 Ops variance due to 2 open headcount positions (hiring freeze)."),
        ("Note:", "Q1 Marketing overspend approved post-hoc — one-time campaign."),
        ("TODO:", "FX adjustment for CC100 international bookings not yet applied."),
        ("TODO:", "Rechnung von IT-Dienstleister (Nov) fehlt noch — erwartet ~18.000 EUR"),
    ]
    for r, (k, v) in enumerate(notes, start=3):
        ws2.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws2.cell(row=r, column=2, value=v)
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 60

    path = DATA_DIR / "budget_tracker.xlsx"
    wb.save(path)
    print(f"✓ Saved {path}")


# ── 2. transactions_q1.csv ───────────────────────────────────────────────────

CATEGORIES = ["Personnel", "Travel", "Software", "External Services",
               "Office Supplies", "Marketing Campaigns", "Training", "Infrastructure"]
CC_SHORT = ["CC100", "CC200", "CC300", "CC400", "CC500"]

def random_date(start: date, end: date) -> str:
    delta = (end - start).days
    return str(start + timedelta(days=random.randint(0, delta)))

def make_transactions_csv():
    rows = []
    start = date(2024, 1, 1)
    end   = date(2024, 3, 31)
    txn_id = 100001
    for _ in range(180):
        cc  = random.choice(CC_SHORT)
        cat = random.choice(CATEGORIES)
        # Marketing gets higher Marketing Campaigns spend; IT gets Software
        if cc == "CC200" and random.random() < 0.5:
            cat = "Marketing Campaigns"
        if cc == "CC400" and random.random() < 0.4:
            cat = "Software"
        amt = round(random.uniform(200, 18000), 2)
        # Inject one clear anomaly: a duplicate booking
        rows.append({
            "transaction_id": f"TXN-{txn_id}",
            "date": random_date(start, end),
            "cost_center": cc,
            "category": cat,
            "vendor": f"Vendor_{random.randint(1, 40):02d}",
            "amount_eur": amt,
            "currency": "EUR",
            "approved": random.choice(["Yes", "Yes", "Yes", "No"]),  # ~25% unapproved
            "description": f"{cat} expense — {cc}",
        })
        txn_id += 1

    # Deliberate duplicate (same amount, same vendor, 1 day apart) — participants find with Bob
    dup = dict(rows[5])
    dup["transaction_id"] = f"TXN-{txn_id}"
    dup["date"] = str(date.fromisoformat(rows[5]["date"]) + timedelta(days=1))
    rows.append(dup)

    path = DATA_DIR / "transactions_q1.csv"
    with open(path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    print(f"✓ Saved {path}  ({len(rows)} rows)")


# ── 3. fx_rates.csv ─────────────────────────────────────────────────────────

def make_fx_rates_csv():
    rows = []
    base_usd = 1.085
    base_gbp = 0.856
    d = date(2024, 1, 1)
    while d <= date(2024, 3, 31):
        if d.weekday() < 5:  # weekdays only
            usd = round(base_usd + random.uniform(-0.015, 0.015), 4)
            gbp = round(base_gbp + random.uniform(-0.010, 0.010), 4)
            rows.append({"date": str(d), "EUR_USD": usd, "EUR_GBP": gbp, "base": "EUR"})
        d += timedelta(days=1)

    path = DATA_DIR / "fx_rates.csv"
    with open(path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["date", "EUR_USD", "EUR_GBP", "base"])
        writer.writeheader()
        writer.writerows(rows)
    print(f"✓ Saved {path}  ({len(rows)} rows)")


# ── Main ─────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    make_budget_xlsx()
    make_transactions_csv()
    make_fx_rates_csv()
    print("\nAll artefacts created in", DATA_DIR)
